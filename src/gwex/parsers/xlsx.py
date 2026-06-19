"""Microsoft Excel (.xlsx) → IR。

pure 層: 与えられた**バイト列**を BytesIO で解析するだけ（パスを開かない）。
方針: `formattedValue` 相当（data_only のキャッシュ値）、1 行目をヘッダ扱い、
結合セルは左上に値・他は空欄（openpyxl が既にそう返すため自然に一致）。

画像: openpyxl は load 時に埋め込み画像を IR に出さないため、同じバイト列を ZIP として
開き、`xl/media/` の実体と drawing の rels 連鎖でシートに関連付けて抽出する。実バイトは
`Document.media`（base64）に集約し、本文には `Image(FetchedRef(id))` を置く。sandbox の
別プロセス境界（Document の JSON 化）をそのまま越えられる。
"""

from __future__ import annotations

import base64
import datetime as _dt
import hashlib
import logging
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET
from io import BytesIO
from typing import Any, Optional

import openpyxl

from gwex.ir import (
    Document,
    FetchedRef,
    Heading,
    Image,
    MediaItem,
    Paragraph,
    Table,
    Text,
)
from gwex.parsers._grids import collapse_empty

logger = logging.getLogger(__name__)

_NS_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# 対象はラスタ主要形式のみ（Bedrock/ブラウザ非対応の EMF/WMF 等は除外）
_MIME_BY_EXT = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "gif": "image/gif",
    "webp": "image/webp",
}
_MAX_ONE_BYTES = 10 * 1024 * 1024  # 1 画像の上限
_MAX_TOTAL_BYTES = 40 * 1024 * 1024  # 合計の上限


def parse(data: bytes) -> Document:
    wb = openpyxl.load_workbook(BytesIO(data), data_only=True)

    # 画像抽出用に同じバイト列を ZIP として開く（失敗してもテキストは出す）
    entries = _zip_entries(data)

    media = _MediaCollector()
    blocks: list[Any] = []
    for ws in wb.worksheets:
        blocks.append(Heading(level=2, inlines=[Text(value=ws.title)]))
        grid = _used_grid(ws)
        if grid:
            ncols = len(grid[0])
            header = [_cell(grid[0][c]) for c in range(ncols)]
            rows = [
                [_cell(row[c] if c < len(row) else None) for c in range(ncols)]
                for row in grid[1:]
            ]
            blocks.append(Table(header=header, rows=rows))
        if entries:
            blocks.extend(_sheet_images(entries, ws.title, media))

    if entries:
        blocks.extend(_orphan_images(entries, media))

    return Document(title=None, source="xlsx", blocks=blocks, media=media.items)


def _used_grid(ws) -> list[list[Any]]:
    """実データのある行・列だけを残した 2 次元配列を返す。

    結合セル由来の空列・行間スペーサ行はレイアウト情報にすぎず LLM には
    ノイズなので、**全体で空の列・行は畳む**（意味を持つ構造だけ残す原則）。
    列を落としても全行に等しく作用するためヘッダ/データの整合は保たれる。
    """
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return collapse_empty(rows, is_empty=lambda v: v is None)


def _cell(value: Any) -> list:
    """セル 1 つを Block 列（[Paragraph]）に。空セルは空 Paragraph。

    NOTE(v1): ハイパーリンクは未対応（空列畳み込みで座標がずれるため）。後続で
    元座標を保持して対応する。
    """
    text = _fmt(value)
    if not text:
        return [Paragraph(inlines=[])]
    return [Paragraph(inlines=[Text(value=text)])]


def _fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    return str(value)


def _zip_entries(data: bytes) -> dict[str, bytes]:
    """xlsx (zip) の全エントリを {名前: バイト} で返す。壊れていれば空 dict。"""
    try:
        with zipfile.ZipFile(BytesIO(data)) as z:
            return {n: z.read(n) for n in z.namelist()}
    except Exception as e:  # noqa: BLE001 - 壊れた zip でも本文抽出は継続
        logger.warning("xlsx: zip を開けず画像抽出をスキップ: %s", e)
        return {}


class _MediaCollector:
    def __init__(self) -> None:
        self.items: list[MediaItem] = []
        self._by_hash: dict[str, int] = {}
        self._next_id = 1
        self._total = 0

    def add(self, raw: bytes, mime: str) -> Optional[int]:
        h = hashlib.sha1(raw).hexdigest()
        if h in self._by_hash:
            return self._by_hash[h]
        if len(raw) > _MAX_ONE_BYTES:
            logger.warning("xlsx: 画像が大きすぎてスキップ (%d bytes)", len(raw))
            return None
        if self._total + len(raw) > _MAX_TOTAL_BYTES:
            logger.warning("xlsx: 画像合計が上限を超えたため以降をスキップ")
            return None
        media_id = self._next_id
        self._next_id += 1
        self._total += len(raw)
        self._by_hash[h] = media_id
        self.items.append(
            MediaItem(id=media_id, mime=mime, data_b64=base64.b64encode(raw).decode("ascii"))
        )
        return media_id

    def seen(self, raw: bytes) -> bool:
        return hashlib.sha1(raw).hexdigest() in self._by_hash


def _mime_for(key: str) -> Optional[str]:
    ext = key.rsplit(".", 1)[-1].lower() if "." in key else ""
    return _MIME_BY_EXT.get(ext)


def _rel_target(rels_xml: str, *, by_id: Optional[str] = None, contains: Optional[str] = None) -> Optional[str]:
    """rels XML から Relationship を1件探して Target を返す。"""
    for m in re.finditer(r"<Relationship\b[^>]*?/?>", rels_xml):
        el = m.group(0)
        if by_id is not None:
            idm = re.search(r'\bId="([^"]+)"', el)
            if not idm or idm.group(1) != by_id:
                continue
        tm = re.search(r'\bTarget="([^"]+)"', el)
        if not tm:
            continue
        target = tm.group(1)
        if contains is not None and contains not in target:
            continue
        return target
    return None


def _sheet_xml_path(entries: dict[str, bytes], sheet_name: str) -> Optional[str]:
    """シート名 → xl/worksheets/sheetN.xml を rels 連鎖で解決。"""
    wbx = entries.get("xl/workbook.xml", b"").decode("utf-8", "replace")
    rid = None
    for m in re.finditer(r"<sheet\b[^>]*?/?>", wbx):
        el = m.group(0)
        nm = re.search(r'\bname="([^"]*)"', el)
        if nm and nm.group(1) == sheet_name:
            ridm = re.search(r'r:id="(rId\d+)"', el)
            rid = ridm.group(1) if ridm else None
            break
    if not rid:
        return None
    rels = entries.get("xl/_rels/workbook.xml.rels", b"").decode("utf-8", "replace")
    target = _rel_target(rels, by_id=rid)
    if not target:
        return None
    return _norm("xl", target)


def _drawing_path(entries: dict[str, bytes], sheet_xml: str) -> Optional[str]:
    base = sheet_xml.split("/")[-1]
    rels = entries.get(f"xl/worksheets/_rels/{base}.rels", b"").decode("utf-8", "replace")
    if not rels:
        return None
    target = _rel_target(rels, contains="drawings/drawing")
    if not target:
        return None
    return _norm(posixpath.dirname(sheet_xml), target)


def _drawing_rels(entries: dict[str, bytes], drawing_path: str) -> dict[str, str]:
    """drawing の rId → media のキー（xl/media/...）。"""
    base = drawing_path.split("/")[-1]
    rels = entries.get(f"xl/drawings/_rels/{base}.rels", b"").decode("utf-8", "replace")
    out: dict[str, str] = {}
    if not rels:
        return out
    for m in re.finditer(r"<Relationship\b[^>]*?/?>", rels):
        el = m.group(0)
        idm = re.search(r'\bId="([^"]+)"', el)
        tm = re.search(r'\bTarget="([^"]+)"', el)
        if idm and tm and "media/" in tm.group(1):
            out[idm.group(1)] = _norm(posixpath.dirname(drawing_path), tm.group(1))
    return out


def _norm(base_dir: str, target: str) -> str:
    """rels の Target（相対 or 絶対）を zip 内の正規キーに直す。"""
    t = target.lstrip("/")
    if t.startswith("xl/"):
        return t
    return posixpath.normpath(posixpath.join(base_dir, target))


def _sheet_images(entries: dict[str, bytes], sheet_name: str, media: _MediaCollector) -> list[Image]:
    sheet_xml = _sheet_xml_path(entries, sheet_name)
    if not sheet_xml:
        return []
    drawing_path = _drawing_path(entries, sheet_xml)
    if not drawing_path or drawing_path not in entries:
        return []
    rels = _drawing_rels(entries, drawing_path)
    try:
        root = ET.fromstring(entries[drawing_path])
    except ET.ParseError as e:
        logger.warning("xlsx: drawing 解析失敗 %s: %s", drawing_path, e)
        return []

    out: list[Image] = []
    for pic in root.iter(f"{{{_NS_XDR}}}pic"):
        blip = pic.find(f".//{{{_NS_A}}}blip")
        embed = blip.get(f"{{{_NS_R}}}embed") if blip is not None else None
        key = rels.get(embed) if embed else None
        if not key or key not in entries:
            continue
        mime = _mime_for(key)
        if not mime:
            logger.warning("xlsx: 非対応形式のためスキップ: %s", key)
            continue
        media_id = media.add(entries[key], mime)
        if media_id is None:
            continue
        out.append(Image(src=FetchedRef(id=media_id, mime=mime), alt=_pic_alt(pic)))
    return out


def _pic_alt(pic: ET.Element) -> Optional[str]:
    cnv = pic.find(f".//{{{_NS_XDR}}}cNvPr")
    if cnv is None:
        return None
    alt = cnv.get("descr") or cnv.get("name")
    return alt or None


def _orphan_images(entries: dict[str, bytes], media: _MediaCollector) -> list[Image]:
    out: list[Image] = []
    for key in sorted(entries):
        if not key.startswith("xl/media/"):
            continue
        mime = _mime_for(key)
        if not mime:
            continue
        raw = entries[key]
        if media.seen(raw):
            continue
        media_id = media.add(raw, mime)
        if media_id is None:
            continue
        out.append(Image(src=FetchedRef(id=media_id, mime=mime), alt=None))
    return out
